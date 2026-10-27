"""Migration of legacy MATLAB-pipeline session metadata into a :class:`SessionConfig`.

Sources (all optional):

- the lab Excel workbook (``CR_Experiments25.xlsx``): one sheet per session with rows
  ``SpikeGLXProbeNr``, ``ProbeArea``, ``ProbeNumber``, ``ProbeDepth``, ``ShankDepth`` and one
  cell per run folder; runs written in *black* font are included, all others excluded
- ``Depths.txt``, ``BrokenChannelsProbe<N>_0Index.txt`` and ``folders2Combine.txt`` in the
  session folder

Values found in the recording files (e.g. probe serial numbers from SpikeGLX meta files)
take precedence; disagreements are reported as warnings.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import numpy as np

from ephysalis.session_config import ProbeConfig, SessionConfig, default_session_config
from ephysalis.session import Session

_RUN_CELL_RE = re.compile(r"^\d+_.+_g\d+$")
_BLACK_RGB = {"FF000000", "00000000"}
_BLACK_THEME = 1  # theme colour "Text 1"
_BLACK_INDEXED = {8, 64}  # palette black, system foreground


@dataclass
class MigrationResult:
    config: SessionConfig
    warnings: list[str] = field(default_factory=list)


def _is_black(cell: Any) -> bool:
    color = cell.font.color if cell.font is not None else None
    if color is None:
        return True
    if color.type == "rgb":
        return color.rgb in _BLACK_RGB
    if color.type == "theme":
        return color.theme == _BLACK_THEME
    if color.type == "indexed":
        return color.indexed in _BLACK_INDEXED
    return True  # "auto"


def _describe_color(cell: Any) -> str:
    color = cell.font.color
    return f"{color.type}={getattr(color, color.type)}"


def _row_values(ws: Any, label: str) -> list[Any] | None:
    """Values to the right of the first cell containing ``label`` (until the first empty cell)."""
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() == label:
                values = []
                col = cell.column + 1
                while (value := ws.cell(row=cell.row, column=col).value) not in (None, ""):
                    values.append(value)
                    col += 1
                return values
    return None


def _number(value: Any) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"-?\d+(\.\d+)?", str(value))
    return float(match.group()) if match else None


def apply_excel_sheet(
    config: SessionConfig, session: Session, workbook_path: Path, sheet_name: str | None = None
) -> list[str]:
    import openpyxl

    warnings: list[str] = []
    wb = openpyxl.load_workbook(workbook_path)
    sheet_name = sheet_name or session.name
    if sheet_name not in wb.sheetnames:
        raise KeyError(f"Sheet {sheet_name!r} not found in {workbook_path}")
    ws = wb[sheet_name]

    # --- probes -----------------------------------------------------------------------
    probe_numbers = _row_values(ws, "SpikeGLXProbeNr") or []
    rows = {
        label: _row_values(ws, label) or []
        for label in ("ProbeArea", "ProbeNumber", "ProbeDepth", "ShankDepth", "ProbeType")
    }
    for i, nr in enumerate(probe_numbers):
        pid = f"imec{int(nr)}"
        probe = config.probes.setdefault(pid, ProbeConfig())

        def get(label: str) -> Any:
            values = rows[label]
            return values[i] if i < len(values) else None

        probe.area = str(get("ProbeArea")) if get("ProbeArea") is not None else probe.area
        probe.insertion_depth_um = _number(get("ProbeDepth")) if get("ProbeDepth") is not None else None
        probe.shank_depth_um = _number(get("ShankDepth")) if get("ShankDepth") is not None else None

        excel_serial = get("ProbeNumber")
        if excel_serial is not None:
            excel_serial = str(int(excel_serial)) if isinstance(excel_serial, (int, float)) else str(excel_serial)
            if probe.serial_number and probe.serial_number != excel_serial:
                warnings.append(
                    f"{pid}: Excel ProbeNumber {excel_serial} differs from serial number "
                    f"{probe.serial_number} in the recording metadata (keeping the recording's)"
                )
            probe.serial_number = probe.serial_number or excel_serial
        if pid not in session.probe_ids:
            warnings.append(f"{pid} listed in Excel but not found in the recording")

    # --- runs ---------------------------------------------------------------------------
    excel_runs: dict[str, tuple[bool, str, str]] = {}
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and _RUN_CELL_RE.match(cell.value.strip()):
                comments = [
                    str(c.value).strip()
                    for c in row
                    if c.column > cell.column and c.value not in (None, "")
                ]
                color = _describe_color(cell) if cell.font.color else ""
                excel_runs[cell.value.strip()] = (_is_black(cell), color, "; ".join(comments))

    for run in config.runs:
        if run.name not in excel_runs:
            run.include = False
            run.notes = "not listed in Excel sheet"
            warnings.append(f"Run {run.name} is not listed in the Excel sheet -> excluded")
            continue
        included, color, run_notes = excel_runs[run.name]
        run.include = included
        run.notes = run_notes if included else f"excluded in Excel (font {color}). {run_notes}".strip()
    for name in excel_runs.keys() - {r.name for r in config.runs}:
        warnings.append(f"Run {name} is listed in the Excel sheet but not found on disk")

    general = [
        str(c.value).strip()
        for row in ws.iter_rows(max_row=12)
        for c in row
        if isinstance(c.value, str) and len(c.value) > 40
    ]
    if general:
        config.notes = "\n".join(general)
    return warnings


def apply_legacy_text_files(
    config: SessionConfig, session: Session, use_folders_file: bool = True
) -> list[str]:
    warnings: list[str] = []
    depths_file = session.path / "Depths.txt"
    if depths_file.exists():
        depths = np.atleast_1d(np.loadtxt(depths_file))
        for i, depth in enumerate(depths):
            pid = f"imec{i}"
            if pid in config.probes and config.probes[pid].insertion_depth_um is None:
                config.probes[pid].insertion_depth_um = float(depth)

    for pid, probe in config.probes.items():
        broken = session.path / f"BrokenChannelsProbe{pid.removeprefix('imec')}_0Index.txt"
        if broken.exists():
            probe.bad_channels = sorted({int(c) for c in np.atleast_1d(np.loadtxt(broken))})

    folders_file = session.path / "folders2Combine.txt"
    if use_folders_file and folders_file.exists():
        listed = [line.strip() for line in folders_file.read_text().splitlines() if line.strip()]
        for run in config.runs:
            if run.name not in listed and run.include:
                run.include = False
                run.notes = (run.notes + " not in folders2Combine.txt").strip()
                warnings.append(f"Run {run.name} not in folders2Combine.txt -> excluded")
    return warnings


def migrate_session(
    session: Session, excel: Path | None = None, sheet: str | None = None
) -> MigrationResult:
    config = default_session_config(session)
    warnings = []
    if excel is not None:
        warnings += apply_excel_sheet(config, session, excel, sheet)
    # as in the MATLAB pipeline, folders2Combine.txt is only used without an Excel sheet
    warnings += apply_legacy_text_files(config, session, use_folders_file=excel is None)
    return MigrationResult(config=config, warnings=warnings)
