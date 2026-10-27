import openpyxl
import pytest
from openpyxl.styles import Font

from ephysalis.session_config import (
    AuxChannelMap,
    SessionConfig,
    default_session_config,
    load_session_config,
    save_session_config,
)
from ephysalis.io.base import open_session
from ephysalis.migrate import migrate_session

RUNS = [("3_rfmapping_g0", "2025-01-09T13:27:45"), ("5_DoseResponse_g0", "2025-01-09T13:58:53")]


def test_default_config_and_roundtrip(make_spikeglx_session, tmp_path):
    session = open_session(make_spikeglx_session(runs=RUNS))
    config = default_session_config(session)
    assert config.subject_id == "SOMxCD1_172"
    assert config.included_runs == ["3_rfmapping_g0", "5_DoseResponse_g0"]
    assert config.probes["imec0"].serial_number == session.runs[0].probes["imec0"].serial_number
    assert config.aux_channels == AuxChannelMap()

    config.probes["imec0"].bad_channels = [3, 191]
    path = tmp_path / "session.toml"
    save_session_config(config, path)
    assert load_session_config(path) == config


def test_newer_schema_is_rejected(tmp_path):
    path = tmp_path / "session.toml"
    path.write_text('schema_version = 99\nsession_id = "x"\nsubject_id = "y"\n')
    with pytest.raises(ValueError, match="newer"):
        load_session_config(path)


def test_legacy_aux_defaults():
    assert AuxChannelMap.legacy_default(8).analog["photodiode"] == 2
    assert AuxChannelMap.legacy_default(4).analog["photodiode"] == 1


def _workbook(path, sheet, serial):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    rows = {
        33: ["SpikeGLXProbeNr", 0],
        34: ["ProbeArea", "V1"],
        35: ["ProbeNumber", serial],
        39: ["ProbeDepth", 2000],
        40: ["ShankDepth", 1794],
    }
    for r, values in rows.items():
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=v)
    ws.cell(row=54, column=1, value="3_rfmapping_g0").font = Font(color="FF000000")
    ws.cell(row=54, column=3, value="centered on 3,7")
    ws.cell(row=71, column=1, value="5_DoseResponse_g0").font = Font(color="FFFF0000")  # red: excluded
    ws.cell(row=90, column=1, value="7_Missing_g0")
    wb.save(path)


def test_excel_migration(make_spikeglx_session, tmp_path):
    session_path = make_spikeglx_session(runs=RUNS)
    (session_path / "BrokenChannelsProbe0_0Index.txt").write_text("12\n5\n")
    session = open_session(session_path)
    meta_serial = session.runs[0].probes["imec0"].serial_number
    excel = tmp_path / "experiments.xlsx"
    _workbook(excel, session.name, serial=12345)

    result = migrate_session(session, excel)
    config: SessionConfig = result.config

    assert config.run_config("3_rfmapping_g0").include
    assert config.run_config("3_rfmapping_g0").notes == "centered on 3,7"
    assert not config.run_config("5_DoseResponse_g0").include
    probe = config.probes["imec0"]
    assert (probe.area, probe.insertion_depth_um, probe.shank_depth_um) == ("V1", 2000, 1794)
    assert probe.serial_number == meta_serial
    assert probe.bad_channels == [5, 12]
    assert any("differs from serial number" in w for w in result.warnings)
    assert any("7_Missing_g0" in w for w in result.warnings)


def test_schema_migration(tmp_path):
    from ephysalis.schema import migrate
    from ephysalis.session_config import SCHEMA_VERSION

    path = tmp_path / "session.toml"
    path.write_text('schema_version = 1\nsession_id = "x"\nsubject_id = "y"\n[[runs]]\nname = "r"\n')
    config = load_session_config(path)
    assert config.schema_version == SCHEMA_VERSION == 6 and config.runs[0].name == "r"
    assert config.events.exclude_roles == ["sync", "wheel"]
    assert config.calibration.run is None and not config.line_noise.enabled
    save_session_config(config, path)
    assert "schema_version = 6" in path.read_text()

    migrations = {1: lambda d: {k: v for k, v in d.items() if k != "a"} | {"b": d["a"]}, 2: lambda d: {**d, "b": d["b"] * 2}}
    assert migrate({"schema_version": 1, "a": 3}, 3, migrations) == {"schema_version": 3, "b": 6}
    assert migrate({"a": 3}, 2, migrations)["b"] == 3  # no version: oldest
    with pytest.raises(ValueError, match="no migration"):
        migrate({"schema_version": 1, "a": 1}, 4, migrations)
    with pytest.raises(ValueError, match="invalid"):
        migrate({"schema_version": "2"}, 2, migrations)
